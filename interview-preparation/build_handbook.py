"""Master Handbook Builder: Generates both DOCX and PDF versions of the Interview Preparation Handbook."""

import sys
from pathlib import Path

# python-docx
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

# reportlab
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, HRFlowable
)
from reportlab.pdfgen import canvas

# Import content chapters
from handbook_data_part1 import CHAPTER_1, CHAPTER_2, CHAPTER_3, CHAPTER_4, CHAPTER_5, CHAPTER_6
from handbook_data_part2 import CHAPTER_7, CHAPTER_8, CHAPTER_9, CHAPTER_10, CHAPTER_11, CHAPTER_12
from handbook_data_part3 import CHAPTER_13, CHAPTER_14, CHAPTER_15, CHAPTER_16, CHAPTER_17

ALL_CHAPTERS = [
    CHAPTER_1, CHAPTER_2, CHAPTER_3, CHAPTER_4, CHAPTER_5, CHAPTER_6,
    CHAPTER_7, CHAPTER_8, CHAPTER_9, CHAPTER_10, CHAPTER_11, CHAPTER_12,
    CHAPTER_13, CHAPTER_14, CHAPTER_15, CHAPTER_16, CHAPTER_17
]

BASE_DIR = Path(__file__).resolve().parent
DOCX_OUT = BASE_DIR / "Autonomous_Web_Agent_Interview_Handbook.docx"
PDF_OUT = BASE_DIR / "Autonomous_Web_Agent_Interview_Handbook.pdf"


# -------------------------------------------------------------
# 1. REPORTLAB NUMBERED CANVAS (PAGE X OF Y)
# -------------------------------------------------------------
class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            canvas.Canvas.showPage(self)
        canvas.Canvas.save(self)

    def draw_page_decorations(self, page_count):
        if self._pageNumber > 1:  # skip cover
            self.saveState()
            self.setFont("Helvetica", 8)
            self.setFillColor(colors.HexColor("#4A5568"))
            
            # Running Header
            self.drawString(45, 755, "Autonomous Web Agent — Complete Technical Interview Preparation Handbook")
            self.setStrokeColor(colors.HexColor("#CBD5E0"))
            self.setLineWidth(0.5)
            self.line(45, 748, 567, 748)
            
            # Running Footer
            self.line(45, 45, 567, 45)
            self.drawString(45, 32, "Candidate: Prajwal Mahesh | Target Role: AI Software Engineer / Python Developer")
            page_text = f"Page {self._pageNumber} of {page_count}"
            self.drawRightString(567, 32, page_text)
            self.restoreState()


# -------------------------------------------------------------
# 2. GENERATE PDF HANDBOOK
# -------------------------------------------------------------
def generate_pdf():
    print(f"Generating PDF handbook at: {PDF_OUT}")
    doc = SimpleDocTemplate(
        str(PDF_OUT),
        pagesize=letter,
        leftMargin=45,
        rightMargin=45,
        topMargin=54,
        bottomMargin=54
    )
    
    styles = getSampleStyleSheet()
    
    # Custom styles
    title_style = ParagraphStyle(
        'CoverTitle',
        fontName='Helvetica-Bold',
        fontSize=24,
        leading=30,
        textColor=colors.HexColor('#1A365D'),
        alignment=1, # Center
        spaceAfter=10
    )
    subtitle_style = ParagraphStyle(
        'CoverSubtitle',
        fontName='Helvetica',
        fontSize=12,
        leading=16,
        textColor=colors.HexColor('#4A5568'),
        alignment=1,
        spaceAfter=25
    )
    h1_style = ParagraphStyle(
        'Heading1_Custom',
        fontName='Helvetica-Bold',
        fontSize=15,
        leading=19,
        textColor=colors.HexColor('#1A365D'),
        spaceBefore=16,
        spaceAfter=8,
        keepWithNext=True
    )
    h2_style = ParagraphStyle(
        'Heading2_Custom',
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=15,
        textColor=colors.HexColor('#2B6CB0'),
        spaceBefore=10,
        spaceAfter=5,
        keepWithNext=True
    )
    body_style = ParagraphStyle(
        'Body_Custom',
        fontName='Helvetica',
        fontSize=9,
        leading=13,
        textColor=colors.HexColor('#2D3748'),
        spaceAfter=6
    )
    code_style = ParagraphStyle(
        'Code_Custom',
        fontName='Courier',
        fontSize=7.5,
        leading=10,
        textColor=colors.HexColor('#1A202C'),
        backColor=colors.HexColor('#F7FAFC'),
        borderColor=colors.HexColor('#E2E8F0'),
        borderWidth=0.5,
        borderPadding=6,
        spaceAfter=8
    )
    meta_style = ParagraphStyle(
        'Meta_Custom',
        fontName='Helvetica',
        fontSize=8.5,
        leading=12,
        textColor=colors.HexColor('#2D3748')
    )
    meta_bold = ParagraphStyle(
        'Meta_Bold',
        fontName='Helvetica-Bold',
        fontSize=8.5,
        leading=12,
        textColor=colors.HexColor('#1A365D')
    )

    story = []

    # COVER PAGE
    story.append(Spacer(1, 40))
    story.append(Paragraph("AUTONOMOUS WEB AGENT", title_style))
    story.append(Paragraph("Complete Architectural, Source Code & Interview Preparation Handbook", subtitle_style))
    story.append(HRFlowable(width="80%", thickness=2, color=colors.HexColor("#2B6CB0"), spaceAfter=25))
    
    meta_table_data = [
        [Paragraph("Candidate Name:", meta_bold), Paragraph("Prajwal Mahesh", meta_style)],
        [Paragraph("Target Positions:", meta_bold), Paragraph("AI Software Engineer / Python Developer / QA Automation Intern", meta_style)],
        [Paragraph("Repository Status:", meta_bold), Paragraph("Original Standalone Repository (isFork: false)", meta_style)],
        [Paragraph("GitHub URL:", meta_bold), Paragraph("https://github.com/prajwal0104/autonomous-web-agent", meta_style)],
        [Paragraph("Local Workspace:", meta_bold), Paragraph("C:\\Users\\Prajwal\\.gemini\\antigravity\\scratch\\autonomous-web-agent", meta_style)],
        [Paragraph("Core Technologies:", meta_bold), Paragraph("Python 3.13, Playwright, Pydantic v2, Typer, Rich, Pytest, LLM APIs", meta_style)],
        [Paragraph("Verification Status:", meta_bold), Paragraph("End-to-End Verified & Tested (4/4 Automated Tests Passed)", meta_style)],
        [Paragraph("Publication Date:", meta_bold), Paragraph("October 2026", meta_style)],
    ]
    meta_table = Table(meta_table_data, colWidths=[150, 360])
    meta_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F7FAFC')),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#CBD5E0')),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ('LEFTPADDING', (0, 0), (-1, -1), 10),
        ('RIGHTPADDING', (0, 0), (-1, -1), 10),
    ]))
    story.append(meta_table)
    story.append(Spacer(1, 30))
    
    overview_box = Paragraph(
        "<b>Notice to Reviewers & Interviewers:</b> This handbook documents the architecture, source code, "
        "engineering decisions, trade-offs, and failure recovery mechanisms implemented in the Autonomous Web Agent project. "
        "All code citations and line numbers correspond to the active verified repository.",
        ParagraphStyle('Notice', fontName='Helvetica-Oblique', fontSize=8, leading=11, textColor=colors.HexColor('#718096'), alignment=1)
    )
    story.append(overview_box)
    story.append(PageBreak())

    # TABLE OF CONTENTS
    story.append(Paragraph("Table of Contents", h1_style))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#CBD5E0"), spaceAfter=12))
    
    for idx, ch in enumerate(ALL_CHAPTERS, 1):
        ch_title = ch["title"]
        story.append(Paragraph(f"<b>{ch_title}</b>", body_style))
    
    story.append(Spacer(1, 15))
    story.append(PageBreak())

    # CHAPTERS
    for ch in ALL_CHAPTERS:
        story.append(Paragraph(ch["title"], h1_style))
        story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#2B6CB0"), spaceAfter=10))
        
        for sec in ch["sections"]:
            story.append(Paragraph(sec["heading"], h2_style))
            
            if "content" in sec and sec["content"]:
                # Split content into paragraphs
                paragraphs = sec["content"].split("\n\n")
                for p_text in paragraphs:
                    if p_text.strip().startswith("```python") or p_text.strip().startswith("```text"):
                        code_content = p_text.replace("```python", "").replace("```text", "").replace("```", "").strip()
                        story.append(Paragraph(code_content.replace("\n", "<br/>").replace(" ", "&nbsp;"), code_style))
                    else:
                        clean_text = p_text.strip().replace("\n", "<br/>")
                        story.append(Paragraph(clean_text, body_style))
            
            if "table" in sec:
                tbl = sec["table"]
                headers = [Paragraph(f"<b>{h}</b>", ParagraphStyle('TH', fontName='Helvetica-Bold', fontSize=8, leading=10, textColor=colors.white)) for h in tbl["headers"]]
                rows = []
                for row in tbl["rows"]:
                    formatted_row = [Paragraph(cell, ParagraphStyle('TD', fontName='Helvetica', fontSize=7.5, leading=9.5, textColor=colors.HexColor('#2D3748'))) for cell in row]
                    rows.append(formatted_row)
                
                pdf_table = Table([headers] + rows, colWidths=[100, 70, 110, 110, 130][:len(headers)])
                pdf_table.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1A365D')),
                    ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E0')),
                    ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
                    ('TOPPADDING', (0, 0), (-1, -1), 4),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
                    ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F7FAFC')])
                ]))
                story.append(Spacer(1, 4))
                story.append(pdf_table)
                story.append(Spacer(1, 6))

        story.append(Spacer(1, 10))

    doc.build(story, canvasmaker=NumberedCanvas)
    print(f"[SUCCESS] PDF generated cleanly at {PDF_OUT}")


# -------------------------------------------------------------
# 3. GENERATE DOCX HANDBOOK
# -------------------------------------------------------------
def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def generate_docx():
    print(f"Generating DOCX handbook at: {DOCX_OUT}")
    doc = docx.Document()

    # Set margins
    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    # COVER PAGE
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("AUTONOMOUS WEB AGENT\n")
    r_title.bold = True
    r_title.font.size = Pt(24)
    r_title.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

    r_sub = p_title.add_run("Complete Architectural, Source Code & Interview Preparation Handbook\n")
    r_sub.font.size = Pt(13)
    r_sub.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)

    doc.add_paragraph()

    meta_table = doc.add_table(rows=8, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Candidate Name:", "Prajwal Mahesh"),
        ("Target Positions:", "AI Software Engineer / Python Developer / QA Automation Intern"),
        ("Repository Status:", "Original Standalone Repository (isFork: false)"),
        ("GitHub URL:", "https://github.com/prajwal0104/autonomous-web-agent"),
        ("Local Workspace:", "C:\\Users\\Prajwal\\.gemini\\antigravity\\scratch\\autonomous-web-agent"),
        ("Core Technologies:", "Python 3.13, Playwright, Pydantic v2, Typer, Rich, Pytest, LLM APIs"),
        ("Verification Status:", "End-to-End Verified & Tested (4/4 Automated Tests Passed)"),
        ("Publication Date:", "October 2026")
    ]
    for i, (k, v) in enumerate(meta_data):
        row = meta_table.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.5)
        r0 = c0.paragraphs[0].add_run(k)
        r0.bold = True
        r0.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
        c1.paragraphs[0].add_run(v)
        set_cell_background(c0, "F7FAFC")
        set_cell_background(c1, "F7FAFC")

    doc.add_page_break()

    # TABLE OF CONTENTS
    h_toc = doc.add_heading("Table of Contents", level=1)
    h_toc.runs[0].font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    for ch in ALL_CHAPTERS:
        p = doc.add_paragraph()
        r = p.add_run(ch["title"])
        r.bold = True

    doc.add_page_break()

    # CHAPTERS
    for ch in ALL_CHAPTERS:
        h1 = doc.add_heading(ch["title"], level=1)
        h1.runs[0].font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

        for sec in ch["sections"]:
            h2 = doc.add_heading(sec["heading"], level=2)
            h2.runs[0].font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)

            if "content" in sec and sec["content"]:
                paragraphs = sec["content"].split("\n\n")
                for p_text in paragraphs:
                    if p_text.strip().startswith("```python") or p_text.strip().startswith("```text"):
                        code_content = p_text.replace("```python", "").replace("```text", "").replace("```", "").strip()
                        p_code = doc.add_paragraph()
                        p_code.paragraph_format.left_indent = Inches(0.3)
                        r_code = p_code.add_run(code_content)
                        r_code.font.name = "Consolas"
                        r_code.font.size = Pt(8.5)
                        r_code.font.color.rgb = RGBColor(0x1A, 0x20, 0x2C)
                    else:
                        doc.add_paragraph(p_text.strip())

            if "table" in sec:
                tbl_data = sec["table"]
                doc_table = doc.add_table(rows=len(tbl_data["rows"]) + 1, cols=len(tbl_data["headers"]))
                doc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                # Header row
                for col_idx, h_text in enumerate(tbl_data["headers"]):
                    cell = doc_table.rows[0].cells[col_idx]
                    run = cell.paragraphs[0].add_run(h_text)
                    run.bold = True
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    set_cell_background(cell, "1A365D")
                # Data rows
                for r_idx, row in enumerate(tbl_data["rows"], 1):
                    bg_color = "F7FAFC" if r_idx % 2 == 0 else "FFFFFF"
                    for col_idx, cell_value in enumerate(row):
                        cell = doc_table.rows[r_idx].cells[col_idx]
                        cell.paragraphs[0].add_run(cell_value)
                        set_cell_background(cell, bg_color)

                doc.add_paragraph()

    doc.save(str(DOCX_OUT))
    print(f"[SUCCESS] DOCX generated cleanly at {DOCX_OUT}")


if __name__ == "__main__":
    generate_docx()
    generate_pdf()
